"""Resume formatting service for eZest DOCX template.

This module provides a small service that:
- Takes extracted resume text.
- Uses the existing LLM-based resume extractor (Gemini by default) to get
  structured data.
- Maps that data into the context expected by the eZest DOCX template.
- Renders a client-ready DOCX using `docxtpl` into the results directory.

Phase 1: DOCX only. PDF export will be added in a later iteration.
"""

import logging
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

from core.config import settings
from services.llm_resume_extractor import create_llm_extractor

logger = logging.getLogger(__name__)

try:  # pragma: no cover - optional dependency in some environments
    from docxtpl import DocxTemplate
    DOCXTPL_AVAILABLE = True
except ImportError:  # pragma: no cover - handled gracefully at runtime
    DOCXTPL_AVAILABLE = False
    logger.warning("docxtpl not available. Install: pip install docxtpl to enable DOCX resume formatting")

try:  # pragma: no cover - optional dependency in some environments
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.shared import Pt

    PYTHON_DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover - handled gracefully at runtime
    Document = None  # type: ignore[assignment]
    CT_Tbl = CT_P = Table = Paragraph = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]
    PYTHON_DOCX_AVAILABLE = False
    logger.warning("python-docx not available. Post-processing of eZest DOCX will be disabled.")


@dataclass
class EzestFormattedResume:
    """Represents a single formatted resume file on disk."""

    candidate_name: str
    original_filename: str
    docx_token: str  # opaque token used in download URL
    docx_path: Path


class EZestResumeFormatter:
    """Format resumes into the eZest DOCX client CV template.

    This service is intentionally self-contained so it can be reused from
    different API flows (standalone Resume Formatter page, candidate views, etc.).
    """

    def __init__(self, template_filename: str = "ezest-updated.docx") -> None:
        self.template_path = Path("templates") / "resume_formatter" / template_filename
        self.results_root = Path(settings.results_dir) / "resume_formatter"
        self.results_root.mkdir(parents=True, exist_ok=True)

        if not self.template_path.exists():
            logger.warning("eZest template not found at %s", self.template_path)

    def format_resume_text(
        self,
        *,
        resume_text: str,
        client_name: str,
        original_filename: str,
        user_id: Optional[str] = None,
        jd_text: Optional[str] = None,
        instructions: Optional[str] = None,
    ) -> EzestFormattedResume:
        """Extract structured data with LLM and render a DOCX into results dir."""

        if not DOCXTPL_AVAILABLE:
            raise RuntimeError("docxtpl is not installed; cannot generate DOCX resumes")

        if not resume_text or len(resume_text.strip()) < 50:
            raise ValueError("Resume text is too short for reliable formatting")

        structured = self._extract_with_llm(resume_text)
        context = self._build_ezest_context(
            structured_data=structured,
            client_name=client_name,
            jd_text=jd_text,
            instructions=instructions,
        )

        candidate_name = structured.get("name") or "Candidate"

        token = self._build_docx_token(
            user_id=user_id,
            candidate_name=candidate_name,
            client_name=client_name,
        )
        output_path = self.results_root / token

        self._render_docx(context=context, output_path=output_path)

        return EzestFormattedResume(
            candidate_name=candidate_name,
            original_filename=original_filename,
            docx_token=token,
            docx_path=output_path,
        )

    # --- LLM extraction -------------------------------------------------

    def _extract_with_llm(self, resume_text: str) -> Dict[str, Any]:
        """Use the shared LLMResumeExtractor (Gemini/OpenAI) to get structured data."""

        provider = "gemini"
        api_key = getattr(settings, "gemini_api_key", None) or os.getenv("GEMINI_API_KEY")
        if not api_key:
            raise RuntimeError("Gemini API key not configured for Resume Formatter")

        extractor = create_llm_extractor(provider=provider, api_key=api_key)
        data = extractor.extract(resume_text)
        data = extractor.validate_extraction(data)
        return data

    # --- Context building for eZest template ----------------------------

    def _build_ezest_context(
        self,
        *,
        structured_data: Dict[str, Any],
        client_name: str,
        jd_text: Optional[str],
        instructions: Optional[str],
    ) -> Dict[str, Any]:
        """Map generic structured resume data into eZest-specific context keys."""

        name = structured_data.get("name") or "Candidate"
        title = self._guess_title(structured_data)

        summary_text: str = structured_data.get("summary") or ""
        summary_bullets = self._split_summary_into_bullets(summary_text) if summary_text else []

        skills = structured_data.get("skills") or []
        skills_rows = self._build_skills_rows(skills)

        experience_raw = structured_data.get("work_experience") or []
        experience = self._build_experience(experience_raw)

        other_projects = structured_data.get("projects") or []
        education = structured_data.get("education") or []
        certifications = structured_data.get("certifications") or []
        certifications_rows = self._build_certifications_rows(certifications)

        context: Dict[str, Any] = {
            "contact_info": {
                "name": name,
                "title": title,
            },
            "summary": summary_text,
            "summary_bullets": summary_bullets,
            # Tools & Technologies
            "skills_rows": skills_rows,
            # Work experience / projects
            "experience": experience,
            "other_projects": other_projects,
            # Education & certifications
            "education": education,
            "certifications_rows": certifications_rows,
            # Extra metadata that can be used in template if needed
            "client_name": client_name,
            "jd_text": jd_text or "",
            "instructions": instructions or "",
        }

        return context

    def _guess_title(self, data: Dict[str, Any]) -> str:
        """Infer a reasonable professional title from experience or summary."""

        # Try explicit title in top work experience
        work = data.get("work_experience") or []
        if work:
            title = work[0].get("title") or ""
            if title and len(title.split()) <= 8:
                return title

        # Fallback: use keywords from summary
        summary = (data.get("summary") or "").lower()
        if "data" in summary and "engineer" in summary:
            return "Data Engineer"
        if "data" in summary and "scientist" in summary:
            return "Data Scientist"
        if "full stack" in summary:
            return "Full Stack Developer"
        if "frontend" in summary or "front-end" in summary:
            return "Frontend Developer"
        if "backend" in summary or "back-end" in summary:
            return "Backend Developer"

        return "IT Professional"

    def _split_summary_into_bullets(self, summary: str, max_bullets: int = 5) -> List[str]:
        """Convert a free-form summary into concise bullet points."""

        # Naive sentence splitting; good enough for MVP
        raw_sentences = [s.strip() for s in summary.replace("\n", " ").split(".")]
        sentences = [s for s in raw_sentences if len(s) > 0]
        bullets = sentences[:max_bullets]
        return bullets

    def _build_skills_rows(self, skills: List[Dict[str, Any]], max_rows: int = 8) -> List[Dict[str, str]]:
        """Build category -> skills rows for the Tools & Technologies section.

        The eZest template expects rows with `left` (category) and `right` (skills).
        """

        if not skills:
            return []

        # First, try to use the LLM-provided "category" field so that
        # headings are dynamic per resume (e.g., "ServiceNow Modules",
        # "Development & Scripting", "Integrations"). Only if everything
        # collapses into a single generic bucket like "Technical" do we
        # fall back to a static skill-type heuristic.

        by_category: Dict[str, List[str]] = {}
        for s in skills:
            name = (s.get("name") or "").strip()
            if not name:
                continue
            category = (s.get("category") or "").strip()
            by_category.setdefault(category, []).append(name)

        # Determine if the categories are actually useful (more than one
        # and not all generic labels like "technical").
        generic_labels = {"", "technical", "technical skills", "skills", "other", "others", "misc", "miscellaneous"}
        non_generic_categories = [
            c for c in by_category.keys() if c and c.lower() not in generic_labels
        ]

        rows: List[Dict[str, str]] = []

        if len(by_category) > 1 and non_generic_categories:
            # Use dynamic categories from the model, preserving their order
            # as much as possible but putting non-generic labels first.
            ordered_cats: List[str] = []
            for c in by_category.keys():
                if c in non_generic_categories and c not in ordered_cats:
                    ordered_cats.append(c)
            for c in by_category.keys():
                if c not in ordered_cats:
                    ordered_cats.append(c)

            for cat in ordered_cats:
                names = by_category.get(cat) or []
                unique_sorted = sorted({n.strip() for n in names if n.strip()})
                if not unique_sorted:
                    continue
                label = cat or "Technical"
                rows.append({
                    "left": label,
                    "right": ", ".join(unique_sorted),
                })

            return rows[:max_rows]

        # Fallback: static heuristic buckets based on skill names so that
        # we still avoid a single "Technical" row.

        all_names: List[str] = []
        for s in skills:
            name = (s.get("name") or "").strip()
            if name:
                all_names.append(name)

        if not all_names:
            return []

        buckets: Dict[str, List[str]] = {
            "Languages": [],
            "Frameworks  Libraries": [],
            "Data Access Technologies": [],
            "Tools  Technologies": [],
            "Databases": [],
        }

        language_keywords = [
            "c#",
            "c++",
            "java",
            "typescript",
            "javascript",
            "python",
            "html",
            "css",
            "sql",
        ]
        db_keywords = [
            "sql server",
            "mysql",
            "postgres",
            "postgresql",
            "oracle",
            "mongodb",
            "sqlite",
            "cosmos db",
        ]
        framework_keywords = [
            "asp.net",
            ".net core",
            "angular",
            "react",
            "vue",
            "rxjs",
            "django",
            "flask",
            "spring",
            "wpf",
        ]
        data_access_keywords = [
            "ado.net",
            "entity framework",
            "ef core",
            "linq",
            "stored procedure",
            "stored procedures",
            "hibernate",
            "dapper",
        ]

        for name in all_names:
            lower = name.lower()

            if any(k in lower for k in language_keywords):
                buckets["Languages"].append(name)
            elif any(k in lower for k in framework_keywords):
                buckets["Frameworks  Libraries"].append(name)
            elif any(k in lower for k in data_access_keywords):
                buckets["Data Access Technologies"].append(name)
            elif any(k in lower for k in db_keywords):
                buckets["Databases"].append(name)
            else:
                buckets["Tools  Technologies"].append(name)

        ordered_labels = [
            "Languages",
            "Frameworks  Libraries",
            "Data Access Technologies",
            "Tools  Technologies",
            "Databases",
        ]

        for label in ordered_labels:
            names = buckets.get(label) or []
            if not names:
                continue
            unique_sorted = sorted({n.strip() for n in names if n.strip()})
            if not unique_sorted:
                continue
            rows.append({
                "left": label,
                "right": ", ".join(unique_sorted),
            })

        # Limit number of rows for layout sanity (though we normally expect <= 5)
        return rows[:max_rows]

    def _build_experience(self, work_experience: List[Dict[str, Any]], max_items: int = 6) -> List[Dict[str, Any]]:
        """Normalize work experience entries for the template."""

        normalized: List[Dict[str, Any]] = []
        for exp in work_experience[:max_items]:
            start = exp.get("start_date") or ""
            end = exp.get("end_date") or ("Present" if exp.get("is_current") else "")

            responsibilities = exp.get("responsibilities") or []
            if isinstance(responsibilities, str):
                responsibilities = [r.strip("- ") for r in responsibilities.split("\n") if r.strip()]

            # Project description may be under different keys; prefer explicit project field
            project_desc = exp.get("project_description") or exp.get("description") or ""

            technologies = exp.get("technologies") or []
            if isinstance(technologies, str):
                technologies = [t.strip() for t in technologies.split(",") if t.strip()]

            normalized.append(
                {
                    "start_date": start,
                    "end_date": end,
                    "project_description": project_desc,
                    "technologies": technologies,
                    "responsibilities": responsibilities,
                }
            )

        return normalized

    def _build_certifications_rows(self, certs: List[Dict[str, Any]], max_items: int = 5) -> List[Dict[str, str]]:
        """Flatten certifications into numbered rows for the table."""

        rows: List[Dict[str, str]] = []
        for idx, c in enumerate(certs[:max_items], start=1):
            name = (c.get("name") or "").strip()
            issuer = (c.get("issuer") or "").strip()
            label = name
            if issuer:
                label = f"{name} – {issuer}" if name else issuer

            rows.append({
                "sno": str(idx),
                "authority": label,
            })

        return rows

    # --- DOCX rendering helpers ----------------------------------------

    def _build_docx_token(
        self,
        user_id: Optional[str],
        candidate_name: Optional[str],
        client_name: Optional[str],
    ) -> str:
        """Build a tokenized DOCX filename that embeds candidate and client.

        Pattern: <user_id>__<candidate-slug>__<client-slug>__<uuid>.docx
        """

        import re
        import uuid

        uid = user_id or "anon"

        cand = (candidate_name or "candidate").strip()
        cand = re.sub(r"[^A-Za-z0-9]+", "-", cand)
        cand = cand.strip("-") or "candidate"

        client = (client_name or "client").strip()
        client = re.sub(r"[^A-Za-z0-9]+", "-", client)
        client = client.strip("-") or "client"

        return f"{uid}__{cand}__{client}__{uuid.uuid4().hex}.docx"

    def _render_docx(self, *, context: Dict[str, Any], output_path: Path) -> None:
        """Render the eZest template using docxtpl into the given path."""

        if not DOCXTPL_AVAILABLE:
            raise RuntimeError("docxtpl is not installed; cannot render DOCX")

        if not self.template_path.exists():
            raise FileNotFoundError(f"eZest template not found at {self.template_path}")

        tpl = DocxTemplate(str(self.template_path))
        tpl.render(context)
        tpl.save(str(output_path))

        # Post-process specific tables (Tools & Technologies, Education, Certifications)
        # to better match the handcrafted eZest layout, without touching other sections.
        try:
            self._post_process_docx(output_path=output_path, context=context)
        except Exception as exc:  # pragma: no cover - defensive logging
            logger.warning("Post-processing of eZest DOCX failed: %s", exc)

        logger.info("Generated eZest DOCX at %s", output_path)

    # --- DOCX post-processing helpers ----------------------------------

    def _post_process_docx(self, *, output_path: Path, context: Dict[str, Any]) -> None:
        """Adjust specific tables after docxtpl rendering.

        - Tighten spacing in Tools & Technologies rows.
        - Ensure Education entries appear on separate rows.
        - Clean up Certifications rows so the table looks compact.
        """

        if not PYTHON_DOCX_AVAILABLE or Document is None:
            return

        doc = Document(str(output_path))

        # Tools & Technologies spacing (identify table after the heading)
        try:
            self._tighten_tools_table(doc)
        except Exception as exc:  # pragma: no cover - defensive logging
            logger.warning("Failed to adjust Tools & Technologies table: %s", exc)

        # Education (separate row per education entry)
        education = context.get("education") or []
        if education:
            try:
                self._populate_education_table(doc, education)
            except Exception as exc:  # pragma: no cover - defensive logging
                logger.warning("Failed to populate Education table: %s", exc)

        # Certifications (compact numbered rows)
        cert_rows = context.get("certifications_rows") or []
        if cert_rows:
            try:
                self._populate_certifications_table(doc, cert_rows)
            except Exception as exc:  # pragma: no cover - defensive logging
                logger.warning("Failed to populate Certifications table: %s", exc)

        doc.save(str(output_path))

    def _iter_block_items(self, doc):
        """Yield paragraphs and tables in document order."""

        if not PYTHON_DOCX_AVAILABLE:
            return []

        body = doc.element.body
        for child in body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, doc)
            elif isinstance(child, CT_Tbl):
                yield Table(child, doc)

    def _tighten_tools_table(self, doc) -> None:
        """Reduce extra vertical spacing between Tools & Technologies rows."""

        if not PYTHON_DOCX_AVAILABLE or Pt is None:
            return

        target_table = None
        seen_heading = False

        for block in self._iter_block_items(doc):
            # Locate the "Tools and Technologies" heading, then pick the first
            # table that appears after it.
            if isinstance(block, Paragraph):
                text = (block.text or "").strip()
                if text == "Tools and Technologies":
                    seen_heading = True
            elif isinstance(block, Table) and seen_heading:
                target_table = block
                break

        if target_table is None:
            return

        for row in target_table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)

    def _find_table_by_headers(self, doc, headers: List[str]):
        """Find a table whose first row matches the given header texts."""

        if not PYTHON_DOCX_AVAILABLE:
            return None

        hdrs_lower = [h.strip().lower() for h in headers]
        for table in doc.tables:
            if len(table.rows) == 0:
                continue
            first = table.rows[0]
            if len(first.cells) < len(headers):
                continue
            row_texts = [first.cells[i].text.strip().lower() for i in range(len(headers))]
            if all(row_texts[i].startswith(hdrs_lower[i]) for i in range(len(headers))):
                return table
        return None

    def _clone_row(self, table, row_idx: int):
        """Clone a row in the table while preserving formatting."""

        from copy import deepcopy

        tr = table.rows[row_idx]._tr
        new_tr = deepcopy(tr)
        table._tbl.append(new_tr)
        return table.rows[-1]

    def _clear_data_rows(self, table, header_rows: int = 1) -> None:
        """Remove all rows after the given number of header rows."""

        while len(table.rows) > header_rows:
            table._tbl.remove(table.rows[-1]._tr)

    def _remove_row(self, table, row_idx: int) -> None:
        """Remove a row from a table by index, with a safe fallback."""

        try:
            row = table.rows[row_idx]
            tbl = table._tbl
            tr = row._tr
            tbl.remove(tr)
        except Exception:
            if 0 <= row_idx < len(table.rows):
                for cell in table.rows[row_idx].cells:
                    cell.text = ""

    def _populate_education_table(self, doc, education: List[Dict[str, Any]]) -> None:
        """Populate the Education Details table with one row per entry."""

        table = self._find_table_by_headers(doc, ["Course", "University", "Year"])
        if table is None:
            return

        # Some versions of the template may accidentally contain an extra trailing
        # column. Normalize to 3 columns to avoid stray cells.
        if len(table.columns) > 3:
            for row in table.rows:
                # Repeatedly remove the last cell until only 3 remain
                while len(row.cells) > 3:
                    row._tr.remove(row.cells[-1]._tc)

        if len(table.rows) > 1:
            header_rows = 2
            template_row_idx = 1
        else:
            table.add_row()
            header_rows = 2
            template_row_idx = 1

        self._clear_data_rows(table, header_rows=header_rows)

        for e in education or []:
            new_row = self._clone_row(table, template_row_idx)
            degree = str(e.get("degree", "") or "").strip()
            institution = str(e.get("institution", "") or "").strip()
            year = str(e.get("graduation_date", "") or "").strip()

            new_row.cells[0].text = degree
            new_row.cells[1].text = institution
            new_row.cells[2].text = year

        if len(table.rows) > 1:
            self._remove_row(table, template_row_idx)

    def _populate_certifications_table(self, doc, cert_rows: List[Dict[str, Any]]) -> None:
        """Populate the Certifications table with compact numbered rows."""

        table = self._find_table_by_headers(doc, ["Sr.No", "University"])
        if table is None:
            return

        if len(table.rows) > 1:
            header_rows = 2
            template_row_idx = 1
        else:
            table.add_row()
            header_rows = 2
            template_row_idx = 1

        self._clear_data_rows(table, header_rows=header_rows)

        for row in cert_rows or []:
            new_row = self._clone_row(table, template_row_idx)
            sno = str(row.get("sno", "") or "").strip()
            authority = str(row.get("authority", "") or "").strip()

            new_row.cells[0].text = sno
            new_row.cells[1].text = authority

        if len(table.rows) > 1:
            self._remove_row(table, template_row_idx)
